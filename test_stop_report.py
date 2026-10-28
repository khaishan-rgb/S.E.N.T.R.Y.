"""Stop-by-stop measured running time report: arithmetic, periods, glitch filters, exports."""
import csv
import io
import json
import unittest
from datetime import datetime

import stop_report as sr

STOPS = [{"seq": 1, "code": "A", "name": "Alpha", "dist": 0.0}, {"seq": 2, "code": "B", "name": "Bravo", "dist": 0.5},
         {"seq": 3, "code": "C", "name": "Charlie", "dist": 1.5}, {"seq": 4, "code": "D", "name": "Delta", "dist": 2.0}]


def ts(day, hh, mm):
    return datetime(2026, 10, day, hh, mm, tzinfo=sr.SGT).timestamp()


def row(i, day, hh, mm, marks, day_type="Weekday"):
    return {"id": i, "date": f"2026-10-{day:02d}", "day_type": day_type, "start_ts": ts(day, hh, mm), "rt_min": (marks[-1] or 0) / 60, "cover": 1.0,
            "marks": json.dumps(marks)}


class Segments(unittest.TestCase):
    def test_stop_to_stop_seconds_and_periods(self):
        trips = sr.clean_trips([row(1, 20, 7, 50, [0, 120, 360, 450]), row(2, 20, 8, 10, [0, 180, 480, 600])], STOPS)
        recs = sr.segment_records(trips, STOPS, period=60, basis="trip")
        self.assertEqual([(r["trip"], r["from_code"], r["to_code"], r["sec"], r["period"]) for r in recs],
                         [(1, "A", "B", 120, "07:00-08:00"), (1, "B", "C", 240, "07:00-08:00"), (1, "C", "D", 90, "07:00-08:00"),
                          (2, "A", "B", 180, "08:00-09:00"), (2, "B", "C", 300, "08:00-09:00"), (2, "C", "D", 120, "08:00-09:00")])
        by_stop = sr.segment_records(trips, STOPS, period=60, basis="stop")
        self.assertEqual(by_stop[2]["period"], "07:00-08:00")                       # trip 1 left C at 07:56
        self.assertEqual(sr.segment_records(trips[:1], STOPS, 60, "stop")[2]["depart"], "07:56:00")

    def test_glitches_and_gaps_are_dropped(self):
        # B->C negative, C->D = 0.5 km in 10 s (180 km/h), a missing mark, and one > 1 h
        trips = sr.clean_trips([row(1, 20, 9, 0, [0, 300, 200, 210]), row(2, 20, 9, 5, [0, None, 400, 4500])], STOPS)
        recs = sr.segment_records(trips, STOPS)
        self.assertEqual([(r["trip"], r["from_code"]) for r in recs], [(1, "A")])

    def test_public_holiday_counts_as_sunday_and_route_change_is_flagged(self):
        trips = sr.clean_trips([row(1, 20, 9, 0, [0, 60, 120, 180, 240])], STOPS, holidays=["2026-10-20"])
        self.assertEqual(trips[0]["day_type"], "Sunday/PH")
        self.assertTrue(trips[0]["route_changed"])
        self.assertEqual(len(trips[0]["marks"]), 4)


class Summary(unittest.TestCase):
    def setUp(self):
        rows = [row(i, 20 + i // 4, 7, 10 * (i % 4), [0, 100 + 10 * i, 300 + 10 * i, 400 + 10 * i]) for i in range(8)]
        self.trips = sr.clean_trips(rows, STOPS)
        self.rep = sr.summarise(self.trips, STOPS, period=30, min_n=5)

    def test_distribution_per_stop_pair(self):
        ab = self.rep["segments"][0]
        mins = sorted((100 + 10 * i) / 60 for i in range(8))
        self.assertEqual(ab["n"], 8)
        self.assertAlmostEqual(ab["median"], round((mins[3] + mins[4]) / 2, 2))
        self.assertAlmostEqual(ab["kmh"], round(0.5 / (ab["median"] / 60), 1))
        self.assertEqual(self.rep["periods"], ["07:00-07:30", "07:30-08:00"])
        self.assertEqual(ab["periods"]["07:00-07:30"]["n"], 6)
        self.assertTrue(ab["periods"]["07:30-08:00"]["low_n"])                     # 2 trips < min_n 5

    def test_cumulative_from_first_stop(self):
        last = self.rep["cumulative"][-1]
        self.assertEqual(last["n"], 8)
        self.assertAlmostEqual(last["median"], round(sorted((400 + 10 * i) / 60 for i in range(8))[3:5][0] / 2 + sorted((400 + 10 * i) / 60 for i in range(8))[3:5][1] / 2, 2))
        self.assertEqual(self.rep["cumulative"][0]["median"], 0)

    def test_excel_has_every_sheet_and_the_numbers(self):
        from openpyxl import load_workbook
        recs = sr.segment_records(self.trips, STOPS, 30)
        wb = load_workbook(io.BytesIO(sr.excel({"service": "14", "direction": 1, "generated": "test"}, self.rep, self.trips, STOPS, recs)))
        self.assertEqual(wb.sheetnames, ["Summary", "Stop-to-stop", "Cumulative", "Trips", "Segment data"])
        ws = wb["Stop-to-stop"]
        self.assertEqual(ws["C2"].value, "A Alpha")
        self.assertEqual(ws["H2"].value, self.rep["segments"][0]["median"])
        self.assertEqual(ws.cell(row=5, column=8).value, "=SUM(H2:H4)")
        self.assertEqual(wb["Segment data"].max_row, len(recs) + 1)
        self.assertEqual(wb["Trips"]["G2"].value, "07:00:00")

    def test_large_segment_sheet_points_to_csv(self):
        from openpyxl import load_workbook
        old = sr.EXCEL_SEGMENT_ROWS
        sr.EXCEL_SEGMENT_ROWS = 3
        try:
            recs = sr.segment_records(self.trips, STOPS, 30)
            wb = load_workbook(io.BytesIO(sr.excel({"service": "14", "direction": 1, "generated": "t"}, self.rep, self.trips, STOPS, recs)))
            self.assertIn("Download CSV", wb["Segment data"]["A2"].value)
        finally:
            sr.EXCEL_SEGMENT_ROWS = old

    def test_csv_has_every_record(self):
        recs = sr.segment_records(self.trips, STOPS, 30)
        text = sr.csv_segments(recs)
        self.assertTrue(text.startswith("﻿"))
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        self.assertEqual(rows[0][:3], ["trip_id", "date", "day_type"])
        self.assertEqual(len(rows) - 1, len(recs))


class Demo(unittest.TestCase):
    def test_demo_report_runs_on_real_stops_and_peaks_are_slower(self):
        rows, stops = sr.demo(days=3)
        rep = sr.summarise(sr.clean_trips(rows, stops), stops, 60)
        self.assertGreater(rep["trips"], 150)
        self.assertEqual(len(rep["segments"]), len(stops) - 1)
        total = lambda p: sum(s["periods"][p]["median"] for s in rep["segments"] if p in s["periods"])
        self.assertGreater(total("08:00-09:00"), total("11:00-12:00"))


if __name__ == "__main__":
    unittest.main()
