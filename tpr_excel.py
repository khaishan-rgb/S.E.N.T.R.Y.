"""Excel export of the zero-history Time Period Report (operator TPR layout + charts). V13.8"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L

F = lambda **k: Font(name="Arial", **k)
_t = Side(style="thin", color="BFBFBF")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)
HDR = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="FCE4D6")
MID = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _hdr(c, text):
    c.value = text
    c.font = F(bold=True, color="FFFFFF", size=9)
    c.fill = HDR
    c.alignment = MID
    c.border = BOX


def _tpr_sheet(wb, r):
    d = r["direction"]
    ws = wb.create_sheet(f"TPR D{d}")
    ws.sheet_view.showGridLines = False
    slots = r["tpr"]["slots"]
    ns = len(slots)
    c0 = 5
    ws["B1"] = "Time Period Report (by ITP Pair) - MODELLED, zero trip history"
    ws["B1"].font = F(bold=True, size=13, color="1F4E78")
    ws["B3"] = f"Svc No : {r['service']}"
    ws["F3"] = f"Dir : {d}"
    ws["I3"] = f"Day Type : {r['day_type']}"
    ws["N3"] = f"Generated : {r['generated']}"
    ws["B4"] = (f"Traffic: {r['sources']['traffic']}  |  Passengers: {r['sources']['passengers']}  |  "
                f"Headway: {r['sources']['headway']}  |  Recommended = P{r['tpr']['pctl']} incl. {r['tpr']['recovery_min']} min recovery")
    for a in ("B3", "F3", "I3", "N3"):
        ws[a].font = F(bold=True, size=10)
    ws["B4"].font = F(italic=True, size=8, color="595959")
    # slot header
    for j, lab in enumerate(["Time Slot", "Start Time", "End Time"]):
        ws.cell(6 + j, 2, lab).font = F(bold=True, size=9)
    for i, s in enumerate(slots):
        a, b = s.split("-")
        _hdr(ws.cell(6, c0 + i), i + 1)
        ws.cell(7, c0 + i, a).font = F(size=9)
        ws.cell(8, c0 + i, b).font = F(size=9)
        for rr in (7, 8):
            ws.cell(rr, c0 + i).alignment = MID
            ws.cell(rr, c0 + i).border = BOX
    pairs = r["tpr"]["pairs"]
    # pair blocks first (values), total block refers to them with formulas
    top = 16
    ws.cell(top - 1, 2, "Start").font = F(bold=True, size=9)
    ws.cell(top - 1, 3, "End").font = F(bold=True, size=9)
    rows_travel, rows_dwell = [], []
    for k, p in enumerate(pairs):
        r0 = top + k * 5
        ws.cell(r0, 1, k + 1).font = F(bold=True)
        ws.cell(r0, 2, p["from"]).font = F(bold=True)
        ws.cell(r0, 3, p["to"]).font = F(bold=True)
        ws.cell(r0 + 3, 2, f"{p['km']} km, {p['stops_between']} stops between").font = F(size=8, italic=True, color="595959")
        for j, lab in enumerate(["Travel time", "Dwell Time", "Prop RunTime"]):
            c = ws.cell(r0 + j, c0 - 1, lab)
            c.font = F(size=8, color="595959")
        for i in range(ns):
            col = L(c0 + i)
            ws.cell(r0, c0 + i, p["travel"][i]).font = F(size=9, color="0000FF")
            ws.cell(r0 + 1, c0 + i, p["dwell"][i]).font = F(size=9, color="0000FF")
            ws.cell(r0 + 2, c0 + i, f"={col}{r0}+{col}{r0 + 1}").font = F(size=9, bold=True)
            for j in range(3):
                x = ws.cell(r0 + j, c0 + i)
                x.alignment = MID
                x.border = BOX
                x.number_format = "0.0"
        rows_travel.append(r0)
        rows_dwell.append(r0 + 1)
    # total block rows 9-13
    labels = ["Total (min)", "Travel time", "Dwell Time", "Prop RunTime", f"Recommended RT (P{r['tpr']['pctl']} + recovery)"]
    for j, lab in enumerate(labels):
        ws.cell(9 + j, 2, lab).font = F(bold=True, size=9)
    for i in range(ns):
        col = L(c0 + i)
        ws.cell(10, c0 + i, "=" + "+".join(f"{col}{x}" for x in rows_travel) if rows_travel else 0)
        ws.cell(11, c0 + i, "=" + "+".join(f"{col}{x}" for x in rows_dwell) if rows_dwell else 0)
        ws.cell(12, c0 + i, f"={col}10+{col}11")
        ws.cell(13, c0 + i, r["tpr"]["total"]["recommended"][i]).font = F(size=9, bold=True, color="0000FF")
        for rr in range(9, 14):
            x = ws.cell(rr, c0 + i)
            x.fill = TOT
            x.border = BOX
            x.alignment = MID
            x.number_format = "0.0"
            if rr != 13:
                x.font = F(size=9, bold=(rr == 12))
    ws.cell(14, 2, "Blue = model output; black = Excel formulas. Recommended RT = planning percentile of the route total, allowing for day-to-day variation, plus recovery.").font = F(size=8, italic=True, color="595959")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    for i in range(ns):
        ws.column_dimensions[L(c0 + i)].width = 6.5
    ws.freeze_panes = ws.cell(9, c0)
    return ws


def _chart_sheet(wb, r):
    d = r["direction"]
    ws = wb.create_sheet(f"Graph D{d}")
    ws.sheet_view.showGridLines = False
    slots = r["tpr"]["slots"]
    pairs = r["tpr"]["pairs"]
    ns = len(slots)
    ws["A1"] = f"Svc {r['service']} D{d} - stop-to-stop running time by time period (min, modelled)"
    ws["A1"].font = F(bold=True, size=12, color="1F4E78")
    _hdr(ws.cell(3, 1), "Stop to stop")
    for i, s in enumerate(slots):
        _hdr(ws.cell(3, 2 + i), s)
    _hdr(ws.cell(3, 2 + ns), "Average")
    _hdr(ws.cell(3, 3 + ns), "Peak")
    for k, p in enumerate(pairs):
        rr = 4 + k
        ws.cell(rr, 1, f"{p['from']} > {p['to']}").font = F(bold=True, size=9)
        for i in range(ns):
            c = ws.cell(rr, 2 + i, p["prop"][i])
            c.font = F(size=9)
            c.alignment = MID
            c.border = BOX
        ws.cell(rr, 2 + ns, round(sum(p["prop"]) / ns, 1)).font = F(size=9, bold=True)
        ws.cell(rr, 3 + ns, max(p["prop"])).font = F(size=9, bold=True)
    last = 3 + len(pairs)
    if pairs:
        ws.conditional_formatting.add(f"B4:{L(1 + ns)}{last}", ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile",
                                                                            mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"))
        lc = LineChart()
        lc.title = "Stop-to-stop running time by time period"
        lc.y_axis.title = "Minutes"
        lc.x_axis.title = "Time period"
        lc.height, lc.width = 10, 26
        lc.add_data(Reference(ws, min_col=1, max_col=1 + ns, min_row=4, max_row=last), from_rows=True, titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=2, max_col=1 + ns, min_row=3))
        lc.x_axis.delete = False
        lc.y_axis.delete = False
        ws.add_chart(lc, f"A{last + 3}")
        bc = BarChart()
        bc.type = "bar"
        bc.title = "Average vs peak, stop to stop"
        bc.y_axis.title = "Minutes"
        bc.height, bc.width = 10, 16
        bc.add_data(Reference(ws, min_col=2 + ns, max_col=3 + ns, min_row=3, max_row=last), titles_from_data=True)
        bc.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
        bc.x_axis.scaling.orientation = "maxMin"
        bc.x_axis.delete = False
        bc.y_axis.delete = False
        ws.add_chart(bc, f"P{last + 3}")
    ws.column_dimensions["A"].width = 18
    for i in range(ns + 2):
        ws.column_dimensions[L(2 + i)].width = 7.5


def _assumptions(wb, reports, assumptions, profile_by_day):
    ws = wb.create_sheet("Method & assumptions")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    rows = [
        ("Model", "Zero-history TPR: no completed bus trips are used."),
        ("Section time", "driving(h) + signals + dwell(h), summed over every stop between the chosen timing points"),
        ("Driving", "Live LTA speed-band time for each stop-to-stop link, rescaled to each hour with the hourly speed profile below; never faster than free-flow"),
        ("Signals", f"{assumptions['junctions_per_km']} junctions per km x {assumptions['signal_s_per_junction']} s"),
        ("Dwell per stop", f"P(stop) x ({assumptions['dwell_base_s']} s base + {assumptions['decel_s']} s deceleration + {assumptions['queue_prob']} x {assumptions['queue_s']} s queue) + {assumptions['dwell_per_pax_s']} s x passengers; P(stop) = 1 - exp(-passengers)"),
        ("Passengers per bus", "DataMall Passenger Volume (tap-in + tap-out) of the stop and hour / days of that day type in the month / services at the stop / buses of this service in the hour"),
        ("Recommended RT", "Planning percentile of the route total, allowing driving to vary about 7.5%, dwell about 18% and recovery about 10% day to day"),
        ("Fallback speed", f"{assumptions['fallback_kmh']} km/h off-peak when live speed bands are unavailable"),
    ]
    for rep in reports:
        for k, v in rep["sources"].items():
            rows.append((f"D{rep['direction']} {k}", v))
    ws["A1"] = "Method & assumptions"
    ws["A1"].font = F(bold=True, size=12, color="1F4E78")
    for i, (a, b) in enumerate(rows):
        ws.cell(3 + i, 1, a).font = F(bold=True, size=9)
        ws.cell(3 + i, 2, b).font = F(size=9)
        ws.cell(3 + i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    r0 = 5 + len(rows)
    ws.cell(r0, 1, "Hourly speed profile (1.00 = weekday off-peak)").font = F(bold=True, size=10)
    _hdr(ws.cell(r0 + 1, 1), "Hour")
    days = list(profile_by_day)
    for j, dname in enumerate(days):
        _hdr(ws.cell(r0 + 1, 2 + j), dname)
    for h in range(24):
        ws.cell(r0 + 2 + h, 1, f"{h:02d}:00").font = F(size=9)
        for j, dname in enumerate(days):
            ws.cell(r0 + 2 + h, 2 + j, profile_by_day[dname][h]).font = F(size=9, color="0000FF")


def workbook(reports, assumptions, profile_by_day):
    wb = Workbook()
    wb.remove(wb.active)
    for rep in reports:
        _tpr_sheet(wb, rep)
    for rep in reports:
        _chart_sheet(wb, rep)
    _assumptions(wb, reports, assumptions, profile_by_day)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
